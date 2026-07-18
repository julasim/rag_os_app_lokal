"""
PDF-Export: Text-Extraktion aus verschiedenen Formaten + fpdf2-Rendering.

Unterstützte Eingabeformate:
  PDF   → Passthrough (Bytes unverändert)
  DOCX  → python-docx Paragraphen
  XLSX  → openpyxl Zellen
  HTML  → BeautifulSoup get_text()
  TXT, MD und alles andere → raw read_text()

Zeichenkodierung: fpdf2 built-in fonts sind Latin-1.
Nicht darstellbare Zeichen (außerhalb Latin-1) werden durch '?' ersetzt.
Deutsche Umlaute (ä/ö/ü/ß) sind in Latin-1 enthalten und bleiben erhalten.
"""
from __future__ import annotations

import asyncio
from pathlib import Path


def _extract_text_sync(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        import pymupdf
        with pymupdf.open(str(path)) as doc:
            return "\n\n".join(page.get_text() for page in doc)

    if suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    if suffix in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            lines.append(f"=== {sheet.title} ===")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append("\t".join(cells))
        return "\n".join(lines)

    if suffix in (".html", ".htm"):
        from bs4 import BeautifulSoup
        return BeautifulSoup(
            path.read_text(errors="replace"), "html.parser"
        ).get_text(separator="\n")

    return path.read_text(errors="replace")


def _build_pdf_bytes(text: str, title: str) -> bytes:
    from fpdf import FPDF

    def _latin1(s: str) -> str:
        return s.encode("latin-1", errors="replace").decode("latin-1")

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    if title:
        pdf.set_font("Helvetica", "B", 13)
        pdf.multi_cell(0, 8, _latin1(title[:200]))
        pdf.ln(3)

    pdf.set_font("Helvetica", size=10)
    pdf.multi_cell(0, 5, _latin1(text))

    return bytes(pdf.output())


async def to_pdf_bytes(path: Path, title: str = "") -> bytes:
    """
    Gibt den Dateiinhalt als PDF-Bytes zurück.
    PDF-Originale werden direkt durchgereicht.
    Alle anderen Formate werden via Text-Extraktion konvertiert.
    """
    if path.suffix.lower() == ".pdf":
        return await asyncio.to_thread(path.read_bytes)

    text = await asyncio.to_thread(_extract_text_sync, path)
    return await asyncio.to_thread(_build_pdf_bytes, text, title or path.name)
